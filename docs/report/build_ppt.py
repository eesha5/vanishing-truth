"""Build the CA-3 presentation on the college template (16:9, 13.333 x 7.5 in).

    python docs/report/build_ppt.py

Starts from CA3_template.pptx (the institute's CA-3 template): its title slide is filled in,
its instruction slides are replaced by content slides in the template's order, and every
content slide keeps the template's look (white, Cambria, bold black titles, slide number
bottom right).  Section order and slide counts follow the template's contents slide.

Writing rules, as for the report: plain sentences, technical terms kept, no em or en dashes
(the build fails if one appears on a slide or in the notes).
"""

import copy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
RES = HERE.parents[1] / "results"
TEMPLATE = HERE / "CA3_template.pptx"
OUT = HERE / "CA3_Presentation_Geometric_Consistency.pptx"

INK = RGBColor(0x00, 0x00, 0x00)
GREY = RGBColor(0x55, 0x55, 0x55)
ACCENT = RGBColor(0xB0, 0x1E, 0x23)        # the red of the institute's logo, used sparingly
LIGHT = RGBColor(0xF2, 0xF2, 0xF2)
PLAN = RGBColor(0xF6, 0xD5, 0xD6)
FONT = "Cambria"
TABLE_STYLE = "{252108C0-F26C-423F-91A2-4CEDFFE9ED8A}"   # the template's own table style
W = 13.333

STUDENTS = [("Aarushi Rudra", "24070127001"), ("Ankur Saxena", "24070127019"),
            ("Arnav Vadhera", "24070127025"), ("Eesha Masand", "24070127043")]
GUIDE = ["Dr. Praween Nishad", "Assistant Professor, Department of Robotics and Automation"]

prs = Presentation(str(TEMPLATE))
template_slides = list(prs.slides)
SLIDE_NUMBER = next(sh._element for sh in template_slides[1].shapes
                    if sh.is_placeholder and sh.placeholder_format.idx == 12)
CONTENT_LAYOUT = template_slides[1].slide_layout


# ------------------------------------------------------------------ helpers
def run_style(r, size, bold=False, color=INK, italic=False):
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = FONT


def slide(title_text, sub=None):
    """A content slide on the template's layout: title placeholder, slide number, no body.
    sub: one line under the title (used to state each result's finding)."""
    s = prs.slides.add_slide(CONTENT_LAYOUT)
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx != 0:
            ph._element.getparent().remove(ph._element)
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Inches(0.92), Inches(0.2), Inches(11.5), Inches(0.95)
    t.text_frame.text = ""
    r = t.text_frame.paragraphs[0].add_run()
    r.text = title_text
    run_style(r, 34, bold=True)
    s.shapes._spTree.append(copy.deepcopy(SLIDE_NUMBER))
    if sub:
        text(s, 0.92, 1.08, 11.5, 0.45, sub, size=19, bold=True, italic=True, color=ACCENT)
    return s


def text(s, x, y, w, h, lines, size=16, bold=False, color=INK, italic=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space_after=6):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, line in enumerate(lines if isinstance(lines, list) else [lines]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        r = p.add_run()
        r.text = line
        run_style(r, size, bold, color, italic)
    return tb


def bullets(s, x, y, w, h, items, size=16, color=INK, space_after=8, marker="•"):
    """items: str, or (str, bold) tuples."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        txt, bold = item if isinstance(item, tuple) else (item, False)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space_after)
        r = p.add_run()
        r.text = f"{marker}  {txt}" if marker else txt
        run_style(r, size, bold, color)
    return tb


def lbullets(s, x, y, w, h, items, size=16, space_after=10):
    """items: (bold label, rest of the sentence), one paragraph each."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (label, rest) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space_after)
        r = p.add_run()
        r.text = f"\u2022  {label} "
        run_style(r, size, bold=True)
        r = p.add_run()
        r.text = rest
        run_style(r, size)
    return tb


def box(s, x, y, w, h, fill=LIGHT, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1.25)
    sh.shadow.inherit = False
    return sh


def heading(s, x, y, w, label, color=ACCENT, size=15):
    text(s, x, y, w, 0.4, label, size=size, bold=True, color=color)


def stat(s, x, y, w, value, label, color=ACCENT):
    text(s, x, y, w, 0.7, value, size=32, bold=True, color=color, align=PP_ALIGN.CENTER)
    text(s, x, y + 0.7, w, 0.7, label, size=12.5, color=GREY, align=PP_ALIGN.CENTER)


def picture(s, fname, x, y, w=None, h=None):
    p = RES / fname
    if not p.exists():
        raise FileNotFoundError(f"results/{fname} is missing; the deck must not ship with a hole")
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return s.shapes.add_picture(str(p), Inches(x), Inches(y), **kw)


def table(s, x, y, w, h, headers, rows, col_w, size=12, shade=None):
    """Template-styled table. shade: {(row, col): RGBColor} for body cells (row 1 = first body row)."""
    shp = s.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = shp.table
    tbl._tbl.tblPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId").text = TABLE_STYLE
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw)
    for i, row in enumerate([headers] + rows):
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.text = ""
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.margin_left = c.margin_right = Inches(0.06)
            for k, ln in enumerate(str(val).splitlines() or [""]):
                p = c.text_frame.paragraphs[0] if k == 0 else c.text_frame.add_paragraph()
                p.alignment = PP_ALIGN.CENTER if i == 0 else PP_ALIGN.LEFT
                r = p.add_run()
                r.text = ln
                run_style(r, size, bold=(i == 0))
            if shade and (i, j) in shade:
                c.fill.solid()
                c.fill.fore_color.rgb = shade[(i, j)]
    return tbl


def notes(s, txt):
    s.notes_slide.notes_text_frame.text = txt


# ================================================================== 1 title (the template's own slide)
title_slide = template_slides[0]
for sh in title_slide.shapes:
    if not sh.has_text_frame:
        continue
    tf = sh.text_frame
    first = tf.text.strip()
    if first == "Title of the Project":
        sh.left, sh.top, sh.width, sh.height = Inches(0.9), Inches(2.45), Inches(11.53), Inches(1.25)
        tf.word_wrap = True
        tf.paragraphs[0].runs[0].text = "Geometric Consistency Analysis for Detection of Synthetic Imagery"
        tf.paragraphs[0].runs[0].font.size = Pt(32)
    elif first == "Name of Students (PRN)":
        sh.left, sh.top, sh.width, sh.height = Inches(0.9), Inches(3.95), Inches(6.0), Inches(2.3)
        tf.word_wrap = True
        tf.paragraphs[0].alignment = PP_ALIGN.LEFT
        tf.paragraphs[0].runs[0].font.size = Pt(20)
        for name, prn in STUDENTS:
            p = tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            r = p.add_run()
            r.text = f"{name} ({prn})"
            run_style(r, 17)
    elif first == "Name and Designation of Guide":
        sh.left, sh.top, sh.width, sh.height = Inches(7.4), Inches(3.95), Inches(5.2), Inches(1.6)
        tf.word_wrap = True
        tf.paragraphs[0].alignment = PP_ALIGN.LEFT
        tf.paragraphs[0].runs[0].font.size = Pt(20)
        for line in GUIDE:
            p = tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            r = p.add_run()
            r.text = line
            run_style(r, 17)
notes(title_slide, "Open with the one-line idea: a real photo is made by one camera, an AI image is not, "
                   "and that difference can be measured in degrees and pixels. Code and data: "
                   "github.com/shmizi/perspective-can-t-lie")

# ================================================================== 2 contents
s = slide("Content of Presentation")
bullets(s, 1.0, 1.4, 11.3, 5.4, [
    ("Introduction", True), ("Review of Literature", True), ("Research Gap", True),
    ("Statement of the Problem", True), ("Objectives of the Study", True),
    ("Methodology, Tools and Techniques", True), ("Results and Discussion", True),
    ("Conclusion and Future Scope", True), ("Schedule of the Work", True), ("References", True),
], size=20, space_after=7)

# ================================================================== 3 introduction 1
s = slide("Introduction: Background and Rationale")
bullets(s, 0.92, 1.3, 11.5, 2.4, [
    "A real photo is made by ONE camera: light passes through one optical centre onto a flat sensor "
    "(the pinhole camera model).",
    "So lines that are parallel in the world meet at one vanishing point, and every direction in the "
    "image implies the same focal length and lens centre.",
    "AI image generators have no camera inside them. They learn what photos look like, not how a "
    "camera forms them. Do their images still obey the camera rules?",
], size=16, space_after=7)
picture(s, "explain_example.png", 2.57, 3.6, w=8.2)
notes(s, "The figure is the whole project in one picture. Lines are coloured by direction. In the real "
         "photo two directions give 649 and 713 px against a true focal length of 675 px. In the ChatGPT "
         "image the directions disagree, so no single camera could have taken it.")

# ================================================================== 4 introduction 2
s = slide("Introduction: Significance and Current Trends")
box(s, 0.92, 1.35, 5.6, 5.0, fill=RGBColor(0xFF, 0xFF, 0xFF), line=LIGHT)
heading(s, 1.15, 1.5, 5.1, "SIGNIFICANCE OF THE STUDY")
lbullets(s, 1.15, 2.0, 5.15, 4.3, [
    ("Forensics.", "Pixel and frequency detectors weaken as generators improve. A broken geometry "
                   "rule stays broken however good the texture is."),
    ("Interpretable.", "Says which rule failed and by how much, in degrees and ratios, not just "
                       "“fake”."),
    ("Robotics.", "Navigation, visual servoing and structure from motion assume projective geometry. "
                  "Synthetic training images must obey it."),
], size=16, space_after=12)
box(s, 6.85, 1.35, 5.6, 5.0)
heading(s, 7.1, 1.5, 5.1, "CURRENT TRENDS")
bullets(s, 7.1, 2.0, 5.1, 4.6, [
    "Sarkar et al. (CVPR 2024): classifiers on geometric features alone detect generated images "
    "(AUC 0.88 to 0.94).",
    "Probing studies: generators store depth and surface normals locally, yet the global camera can "
    "still be wrong.",
    "New generation methods force vanishing points to agree, so perspective errors are a known "
    "weakness.",
    "2026: claims that the newest models have fixed geometry. Nobody had tested this. We test it on "
    "Gemini and ChatGPT.",
], size=15.5, space_after=10)

# ================================================================== 5, 6 literature (template table)
LIT_HEAD = ["Sr. No", "Author Name", "Paper Title", "Journal Name", "Methodology", "Remarks (Research Gaps)"]
LIT_W = [0.62, 1.85, 2.95, 1.6, 2.45, 2.62]
LIT = [
    ["1", "A. Sarkar, H. Mai, A. Mahapatra, S. Lazebnik, D. A. Forsyth, A. Bhattad",
     "Shadows Don't Lie and Lines Can't Bend! Generative Models don't know Projective Geometry...for now",
     "CVPR, 2024", "Classifiers on line segments, perspective fields and object-shadow pairs",
     "Detects that geometry is wrong (AUC 0.88 to 0.94) but not which rule fails or by how much"],
    ["2", "E. Kee, J. F. O'Brien, H. Farid", "Exposing photo manipulation with inconsistent shadows",
     "ACM Trans. on Graphics, 2013", "Shadow-object wedge constraints solved as a linear program",
     "Shadow-object pairs marked by hand; not applied to AI images at scale"],
    ["3", "H. Farid", "Perspective (in)consistency of paint by text", "arXiv, 2022",
     "Manual vanishing-point analysis of early text-to-image outputs",
     "Small hand-checked sample; no statistics against many real photos"],
    ["4", "A. Criminisi, I. Reid, A. Zisserman", "Single view metrology", "IJCV, 2000",
     "Measurement from vanishing points and the vanishing line",
     "Assumes a real camera; the basis of our test, not a detector"],
    ["5", "J. M. Coughlan, A. L. Yuille", "Manhattan world: compass direction from a single image by "
     "Bayesian inference", "ICCV, 1999", "Bayesian estimate of three directions at right angles",
     "The right-angle assumption fails on angled streets"],
    ["6", "G. Schindler, F. Dellaert", "Atlanta world: an EM framework for simultaneous low-level edge "
     "grouping and camera calibration", "CVPR, 2004", "One vertical and several horizontal directions, "
     "fitted by EM", "Used for calibration; never used to test whether an image has one camera"],
    ["7", "P. Denis, J. H. Elder, F. J. Estrada", "Efficient edge-based methods for estimating Manhattan "
     "frames in urban imagery", "ECCV, 2008", "York Urban Database: 102 calibrated photos with true "
     "vanishing points", "One camera, hand-picked scenes; a single reference can bias conclusions"],
    ["8", "M. El Banani et al.", "Probing the 3D awareness of visual foundation models", "CVPR, 2024",
     "Linear probes read depth and surface normals from model features",
     "Local 3D only; global camera consistency untested"],
]
for part, rows in enumerate([LIT[:4], LIT[4:]]):
    s = slide(f"Review of Literature ({part + 1}/2)")
    table(s, 0.62, 1.3, sum(LIT_W), 5.0 if part == 0 else 4.6, LIT_HEAD, rows, LIT_W, size=12)
    if part == 1:
        text(s, 0.62, 6.15, 12.1, 0.6,
             "Summary: generated images break projective geometry. Open: which rule, by how much, and "
             "measured against which real photographs.", size=15, italic=True, color=ACCENT)
notes(s, "Sarkar et al. is the paper we extend: from detection to diagnosis. Full list of 26 references "
         "in the report.")

# ================================================================== 7 research gap
s = slide("Research Gap")
GAPS = [("1", "Detection, not diagnosis", "Learned classifiers say a violation exists, not which rule "
                                          "failed or by how much."),
        ("2", "One real reference", "Studies compare against one hand-picked dataset. Whether the "
                                    "conclusion depends on that choice is unexamined."),
        ("3", "Orientation vs camera never separated", "No study separates getting the scene's "
                                                       "directions wrong from getting the camera wrong."),
        ("4", "Applicability ignored", "The three-vanishing-point test does not apply to scenes that "
                                       "are not box-shaped, yet it is applied to them anyway.")]
for i, (n, hd, body) in enumerate(GAPS):
    x, y = 0.92 + (i % 2) * 5.85, 1.4 + (i // 2) * 2.45
    box(s, x, y, 5.6, 2.2)
    text(s, x + 0.25, y + 0.2, 0.6, 0.6, n, size=30, bold=True, color=ACCENT)
    text(s, x + 0.9, y + 0.27, 4.5, 0.5, hd, size=17, bold=True)
    text(s, x + 0.9, y + 0.85, 4.5, 1.25, body, size=15, color=GREY)
text(s, 0.92, 6.4, 11.5, 0.45, "This project addresses gaps 1 to 3, and solves gap 4 as part of the method.",
     size=15, italic=True, color=ACCENT)
notes(s, "Gap 4 turned out to change the results the most: see Results 2.")

# ================================================================== 8 problem statement
s = slide("Statement of the Problem")
box(s, 0.92, 1.4, 11.5, 1.55, line=ACCENT, fill=RGBColor(0xFF, 0xFF, 0xFF))
text(s, 1.2, 1.55, 10.95, 1.3,
     "Given a single image, decide whether its geometry fits ONE pinhole camera. Report any mismatch as a "
     "readable residual tied to a specific geometric rule, compare it with real photographs, and use the "
     "result to compare image generators.", size=18, anchor=MSO_ANCHOR.MIDDLE)
heading(s, 0.92, 3.25, 11.5, "RESEARCH QUESTIONS")
bullets(s, 0.92, 3.75, 11.5, 3.0, [
    "Q1. Do images from current generators obey a single pinhole camera, compared with real photos of "
    "similar scenes?",
    "Q2. If not, what fails: the right angles between directions, or the focal length and lens centre "
    "agreeing across the image?",
    "Q3. Does the problem shrink as models get bigger, and do Gemini and ChatGPT still show it?",
    "Q4. Can simple, readable residuals compete with learned classifiers at telling real from generated?",
], size=18, space_after=12, marker="")

# ================================================================== 9 objectives
s = slide("Objectives of the Study")
OBJ = ["To build and validate a measurement pipeline that tests whether an image's lines are consistent "
       "with a single pinhole camera, using synthetic scenes and real photographs with known cameras.",
       "To compare images from five generators (Stable Diffusion 1.5, SDXL, Flux, Gemini and ChatGPT) with two "
       "sets of real photographs under matched content, and find which camera rule they break and by how "
       "much.",
       "To build an interactive application that measures any uploaded image, explains the result and "
       "gives a calibrated likelihood that the image is AI-generated."]
for i, o in enumerate(OBJ):
    y = 1.45 + i * 1.75
    text(s, 0.92, y, 0.8, 0.8, str(i + 1), size=40, bold=True, color=ACCENT)
    text(s, 1.85, y + 0.12, 10.55, 1.5, o, size=19)
notes(s, "The report lists seven detailed objectives; they group into these three. 1: O1, O2, O4, O6. "
         "2: O3, O5. 3: O7.")

# ================================================================== 10 to 13 methodology
s = slide("Methodology: Objective-wise Flowchart")
picture(s, "objective_flowchart.png", 0.92, 1.2, w=11.5)
notes(s, "Each column is one objective, read top to bottom. The pipeline built for objective 1 is reused "
         "unchanged by objectives 2 and 3, which is why the app and the report always agree.")

s = slide("Methodology: Techniques")
table(s, 0.92, 1.3, 11.5, 2.15, ["Level", "Camera rule", "Residual (units)", "Role"],
      [["L2", "Lines parallel in 3D meet at one vanishing point", "Angle (deg)", "Built, validated"],
       ["L3", "Three directions at right angles give one focal length and lens centre",
        "Deviation from 90 deg", "Secondary"],
       ["L3b (Atlanta)", "Every horizontal direction gives the SAME focal length with the vertical",
        "Spread of log focal length", "PRIMARY measure"],
       ["L7", "All cast shadows agree with one light source", "LP violation (px)", "Synthetic validation"]],
      [1.7, 5.6, 2.4, 1.8], size=13)
box(s, 0.92, 3.7, 5.6, 3.05)
heading(s, 1.15, 3.85, 5.2, "FOCAL LENGTH FROM TWO DIRECTIONS")
text(s, 1.15, 4.3, 5.2, 0.5, "f² = −(v₁ − p) · (v₂ − p)",
     size=22, bold=True)
text(s, 1.15, 4.95, 5.2, 1.75,
     ["v₁, v₂: vanishing points of two perpendicular directions; p: the image centre.",
      "In a real photo every (vertical, horizontal) pair gives the same focal length. A pair with "
      "f² ≤ 0 is one no camera could produce."],
     size=15, space_after=8)
box(s, 6.82, 3.7, 5.6, 3.05)
heading(s, 7.05, 3.85, 5.2, "APPLICABILITY RULE (EVIDENCE ONLY)")
bullets(s, 7.05, 4.35, 5.2, 2.35, [
    "3 line families, each at least 8% of the line length",
    "each vanishing point pinned to under 1 deg (bootstrap)",
    "pairs at least 10 deg apart, and 5 times their uncertainty",
    "never uses the residual, so it cannot bias the test",
    "LSD lines, sequential RANSAC, NO right-angle prior",
], size=15, space_after=6)
notes(s, "Key design decision: the vanishing-point finder is never told to look for right angles. The "
         "camera test asks whether the directions are at right angles, so building that in would make "
         "the test meaningless.")

s = slide("Methodology: Data Acquisition and Tools")
table(s, 0.92, 1.3, 7.55, 4.3, ["Dataset", "Images", "Role"],
      [["York Urban", "102", "Hand-picked real photos with true vanishing points"],
       ["Wikimedia Commons", "358", "Ordinary real photos, 135 camera models, EXIF focal length"],
       ["Stable Diffusion 1.5", "450", "Older open model, run on our GPU, settings logged"],
       ["SDXL", "650", "Newer open model, run on our GPU, settings logged"],
       ["FLUX.1-schnell", "250", "Newest open model (12B, 4-bit), run on our GPU"],
       ["Gemini (“nano banana”)", "181", "Closed frontier model, September 2026"],
       ["ChatGPT image model", "51", "Closed frontier model, September 2026"],
       ["Synthetic scenes", "on demand", "Known cameras and injected violations, for testing"]],
      [2.45, 1.15, 3.95], size=12.5)
text(s, 0.92, 5.8, 7.55, 1.0,
     "Three prompt sets with content matched by construction. Every comparison stays within one prompt "
     "set, so a model is never compared on different scenes.", size=14, italic=True, color=GREY)
box(s, 8.75, 1.3, 3.67, 5.45)
heading(s, 8.95, 1.45, 3.3, "TOOLS")
bullets(s, 8.95, 1.95, 3.3, 2.6, [
    "Python 3.14, OpenCV 4.13 (LSD)",
    "NumPy, SciPy (least squares, LP)",
    "scikit-learn, statsmodels",
    "PyTorch 2.14, diffusers 0.40, GGUF",
    "Streamlit 1.58 (demo app)",
    "RTX 5060 Laptop GPU, 8 GB",
], size=14, space_after=5)
heading(s, 8.95, 4.45, 3.3, "STATISTICS")
bullets(s, 8.95, 4.95, 3.3, 1.65, [
    "Bootstrap 95% CI (2000 resamples)",
    "Wilson intervals, Mann-Whitney",
    "5-fold cross-validation",
], size=14, space_after=5)

s = slide("Methodology: Demonstration App Workflow")
picture(s, "app_workflow_slide.png", 0.67, 1.2, w=12.0)
text(s, 0.92, 6.3, 11.5, 0.5,
     "Streamlit app on the same per-image code as the report. The model sees only the ten measurements; "
     "the built-in examples were held out of training.", size=14, italic=True, color=GREY)
notes(s, "The diamond is the applicability rule: an image without lines in three directions gets no score "
         "instead of a guess. Calibration turns the vote into a probability, re-based so that real and AI "
         "were equally likely before looking.")

# ================================================================== 14 to 17 results
s = slide("Results and Discussion (1/4)", "The measurement recovers real cameras")
for i, (v, lab) in enumerate([("0.3°", "vanishing-point error, synthetic scenes"),
                              ("< 1%", "focal-length error, synthetic scenes"),
                              ("81 / 94%", "York Urban true vanishing points found within 2° / 5°"),
                              ("1.02", "fitted / EXIF focal length on real photos (IQR 0.99 to 1.05)")]):
    stat(s, 0.92 + i * 2.9, 1.6, 2.75, v, lab)
picture(s, "commons_focal_validation.png", 1.3, 3.25, h=3.55)
box(s, 5.6, 3.25, 6.82, 3.55)
bullets(s, 5.85, 3.45, 6.35, 3.25, [
    ("The pipeline recovers the real camera from real photos.", True),
    "Left: fitted focal length against the EXIF focal length for photos from 135 camera models. Dashed "
    "lines: a 20% band.",
    "This has to hold before anything the tool says about AI images can be trusted.",
    ("Objective 1 achieved.", True),
], size=15, space_after=10)

s = slide("Results and Discussion (2/4)", "Stable Diffusion breaks the single camera; Flux comes close")
picture(s, "atlanta_dots_slide.png", 0.55, 1.65, h=4.3)
text(s, 0.6, 6.1, 7.4, 0.6,
     "Line-rich prompt set. Dots: medians; lines: 95% intervals; shaded band: range of the real-photo "
     "intervals.", size=12, italic=True, color=GREY)
box(s, 8.25, 1.7, 4.17, 5.1)
bullets(s, 8.45, 1.85, 3.8, 4.9, [
    ("Real photos: directions agree on one focal length within about 15% (0.14, 0.15).", True),
    ("SD 1.5 and SDXL: 35 to 60% disagreement (0.47, 0.36), p ≈ 3×10⁻⁵.", True),
    ("Flux: 0.18. Not distinguishable from real photos (p ≈ 0.17) and better than both SD models "
     "(p ≤ 0.004). It still has more impossible pairs (48%).", True),
    "Two very different real sets agree, and every result survives both robustness checks.",
    "A naive right-angle test made real street photos look WORSE than AI (12.2°). The "
    "applicability rule fixes this.",
], size=13, space_after=6)
notes(s, "This is the slide to defend. Bootstrap intervals, 2000 resamples; Wilson intervals for "
         "proportions; Mann-Whitney test. SD 1.5 to SDXL improves orientation (2.93 to 1.94 deg) but not "
         "camera coherence (p = 0.55). Flux (250 images, 78% measurable) has a median of 0.18 [0.14, 0.30]; "
         "it stays indistinguishable from real photos with the vertical guard and on tilted cameras only. "
         "Not distinguishable is not the same as proven equal: the upper end of its interval is 0.30.")

s = slide("Results and Discussion (3/4)", "The newest models, Gemini and ChatGPT, still fail")
picture(s, "frontier_dots_slide.png", 0.6, 1.65, h=4.55)
box(s, 8.05, 1.65, 4.37, 4.55)
heading(s, 8.25, 1.8, 4.0, "GEMINI AND CHATGPT")
bullets(s, 8.25, 2.3, 4.0, 3.85, [
    ("Both still fail: p ≈ 0.015 against both real sets.", True),
    "Not yet shown to beat SD 1.5 or SDXL (p = 0.14 to 0.26).",
    "Small samples: 31 and 29 usable images.",
    "ChatGPT's impossible-pair rate (39%) looks real; its focal spread does not.",
    "Made from the first prompt set, so not directly comparable with Flux.",
], size=14, space_after=8)
box(s, 0.6, 6.35, 11.82, 0.5, fill=ACCENT)
text(s, 0.8, 6.35, 11.4, 0.5, "For Gemini and ChatGPT, the claim that the geometric giveaway has "
                              "disappeared does not hold on this measure.", size=15, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF),
     anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Gemini and ChatGPT images were made in September 2026 from the first prompt set, so they are "
         "compared with the SD and SDXL images from that same set.")

s = slide("Results and Discussion (4/4)", "Geometry alone detects AI images, and the app explains why")
stat(s, 0.92, 1.6, 3.7, "0.92", "best AUC, residuals only (learned features in Sarkar et al.: 0.88 to 0.94)")
stat(s, 4.82, 1.6, 3.7, "0.77", "demo app, cross-validated AUC, five generators")
stat(s, 8.72, 1.6, 3.7, "78% / 62%", "AI images flagged / real photos cleared at 50%")
heading(s, 0.92, 3.1, 5.6, "SIX CLASSIFIERS ON RESIDUALS (BEST: RANDOM FOREST)", size=13)
table(s, 0.92, 3.5, 5.6, 2.2, ["Task", "AUC"],
      [["York Urban vs SD 1.5", "0.92"], ["York Urban vs SDXL", "0.88"],
       ["Commons vs SDXL", "0.81"], ["Field of view matched", "0.81"]],
      [4.3, 1.3], size=13)
heading(s, 6.82, 3.1, 5.6, "APP SCORE BY GENERATOR", size=13)
table(s, 6.82, 3.5, 5.6, 2.3, ["Generator", "Seen in training", "Never seen"],
      [["SD 1.5", "0.83", "0.81"], ["SDXL", "0.80", "0.76"], ["Flux", "0.73", "0.64"],
       ["Gemini", "0.74", "0.73"], ["ChatGPT", "0.62", "0.63"]],
      [1.9, 1.85, 1.85], size=12.5)
text(s, 0.92, 5.9, 11.5, 0.95,
     ["Readable residuals reach the accuracy of learned geometric features. Image shape is left out on "
      "purpose: every AI image here is 4:3, so it would be a dataset shortcut, not geometry.",
      "Objectives 2 and 3 achieved. Live demo: One Camera or Not?"],
     size=14, color=GREY, space_after=4)
notes(s, "Six algorithms: logistic regression, decision tree, SVM, random forest, naive Bayes, kNN; "
         "five-fold cross-validation. Adding Flux lowered the app's AUC from 0.80 to 0.77: it is the "
         "hardest generator, which is the honest result. Other classifiers score 0.74 to 0.75. Then "
         "switch to the app: python -m streamlit run app/streamlit_app.py, localhost:8501/?examples=all.")

# ================================================================== 18 conclusion
s = slide("Conclusion and Future Scope")
heading(s, 0.92, 1.3, 6.6, "CONCLUSION")
bullets(s, 0.92, 1.8, 6.6, 5.0, [
    "Stable Diffusion images get the directions roughly right but do not commit to one camera: focal "
    "agreement is 2.5 to 3 times worse than real photos.",
    "A bigger model of the same kind (SD 1.5 to SDXL) improves orientation, not camera coherence.",
    "Flux, a newer architecture, comes close to real photos on focal agreement, though it still makes "
    "more impossible pairs.",
    "Gemini and ChatGPT still fail on the first prompt set.",
    "Results hold for two different real photo sets and survive two robustness checks.",
    "Readable residuals reach learned-detector accuracy, and the app explains every score.",
], size=14.5, space_after=7)
text(s, 0.92, 6.3, 6.6, 0.45, "All three objectives were achieved.", size=16, bold=True, color=ACCENT)
box(s, 7.8, 1.3, 4.62, 5.45)
heading(s, 8.0, 1.45, 4.25, "FUTURE SCOPE")
bullets(s, 8.0, 1.95, 4.25, 4.7, [
    "Line-rich test of Gemini and ChatGPT, plus Midjourney; find out what Flux does differently.",
    "Find WHERE an image breaks: per-region camera checks for inpainting and splicing.",
    "Shadows on real photos, pairing shadows with objects (SAM 2).",
    "Combine with a learned detector; test cropping and compression.",
], size=15, space_after=12)

# ================================================================== 19 schedule
s = slide("Schedule of the Work")
WEEKS = ["Week 1\n14 to 20 Sep", "Week 2\n21 to 27 Sep", "Week 3\n28 Sep to 4 Oct", "Week 4\nAnd ahead"]
TASKS = [("Literature review, problem statement, plan", {0}, "Done"),
         ("Geometry pipeline and synthetic validation (Obj. 1)", {0}, "Done"),
         ("Real-photo baselines: York Urban, Commons (Obj. 1)", {0, 1}, "Done"),
         ("Generated sets: SD 1.5, SDXL (Obj. 2)", {0, 1, 2}, "Done"),
         ("Applicability rule, Atlanta measure, classifiers (Obj. 1, 2)", {1, 2}, "Done"),
         ("Frontier models: Gemini, ChatGPT (Obj. 2)", {2}, "Done"),
         ("Robustness checks and demo app (Obj. 2, 3)", {2}, "Done"),
         ("Report, slides, code release on GitHub", {2}, "Done"),
         ("Flux: 250 images, measured, app retrained (Obj. 2, 3)", {2}, "Done"),
         ("Future scope: more frontier and real images for training", {3}, "Planned")]
rows, shade = [], {}
for i, (task, weeks, status) in enumerate(TASKS, start=1):
    rows.append([task] + ["" for _ in WEEKS] + [status])
    for w in weeks:
        shade[(i, w + 1)] = PLAN if status == "Planned" else RGBColor(0x9E, 0x9E, 0x9E)
table(s, 0.62, 1.3, 12.1, 5.2, ["Task"] + WEEKS + ["Status"], rows, [5.0, 1.55, 1.55, 1.55, 1.55, 0.9],
      size=13, shade=shade)
text(s, 0.62, 6.7, 12.1, 0.35, "Grey: completed. Pink: planned.", size=12, italic=True, color=GREY)
notes(s, "Dates follow the project's commit history. Adjust week 1 if the literature review started "
         "earlier.")

# ================================================================== 20 references
s = slide("References")
REFS = [
    "[1] A. Sarkar et al., “Shadows Don’t Lie and Lines Can’t Bend! Generative Models don’t "
    "know Projective Geometry…for now,” CVPR, 2024.",
    "[2] R. Hartley and A. Zisserman, Multiple View Geometry in Computer Vision, 2nd ed., Cambridge "
    "University Press, 2004.",
    "[3] E. Kee, J. F. O’Brien and H. Farid, “Exposing photo manipulation with inconsistent "
    "shadows,” ACM TOG, vol. 32, no. 3, 2013.",
    "[4] H. Farid, “Perspective (in)consistency of paint by text,” arXiv, 2022.",
    "[5] A. Criminisi, I. Reid and A. Zisserman, “Single view metrology,” IJCV, vol. 40, no. 2, "
    "pp. 123-148, 2000.",
    "[6] J. M. Coughlan and A. L. Yuille, “Manhattan world,” ICCV, 1999.",
    "[7] G. Schindler and F. Dellaert, “Atlanta world,” CVPR, 2004.",
    "[8] P. Denis, J. H. Elder and F. J. Estrada, “Efficient edge-based methods for estimating "
    "Manhattan frames in urban imagery,” ECCV, 2008.",
    "[9] R. Grompone von Gioi et al., “LSD: a fast line segment detector with a false detection "
    "control,” IEEE TPAMI, vol. 32, no. 4, 2010.",
    "[10] M. El Banani et al., “Probing the 3D awareness of visual foundation models,” CVPR, 2024.",
    "[11] D. Podell et al., “SDXL: improving latent diffusion models for high-resolution image "
    "synthesis,” ICLR, 2024.",
    "[12] L. Breiman, “Random forests,” Machine Learning, vol. 45, no. 1, 2001.",
]
text(s, 0.92, 1.25, 11.5, 5.3, REFS, size=15, space_after=5)
text(s, 0.92, 6.6, 11.5, 0.35, "Full list of 26 references in the report. Datasets: York Urban (Elder "
                               "Laboratory, York University); Wikimedia Commons photos under CC licences.",
     size=11, italic=True, color=GREY)

# ------------------------------------------------------------------ drop the template's instruction slides
sld_ids = prs.slides._sldIdLst
for sld_id in list(sld_ids)[1:len(template_slides)]:
    prs.part.drop_rel(sld_id.rId)
    sld_ids.remove(sld_id)


# ------------------------------------------------------------------ checks + save
def _texts(presentation):
    for sl in presentation.slides:
        for shp in sl.shapes:
            if shp.has_text_frame:
                yield shp.text_frame.text
            if shp.has_table:
                for row in shp.table.rows:
                    for c in row.cells:
                        yield c.text
        if sl.has_notes_slide:
            yield sl.notes_slide.notes_text_frame.text


for txt in _texts(prs):
    for bad in ("—", "–"):
        if bad in txt:
            raise ValueError(f"dash {bad!r} found in: {txt[:120]}")

try:
    prs.save(str(OUT))
    saved = OUT
except PermissionError:
    saved = OUT.with_name(OUT.stem + "_v2" + OUT.suffix)
    prs.save(str(saved))
print(f"written: {saved} | slides: {len(prs.slides)}")
